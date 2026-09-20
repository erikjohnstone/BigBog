from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from bactalk.intake import (
    IntakeError,
    parse_bacnet_scan_json,
    parse_points_csv,
    parse_points_file,
    parse_sequence_document,
)

EXAMPLES = Path(__file__).parents[1] / "examples"


def test_contractor_csv_and_scan_parse_and_cross_validate() -> None:
    points = parse_points_csv((EXAMPLES / "vav-reheat-points.csv").read_bytes())
    scan = parse_bacnet_scan_json((EXAMPLES / "vav-reheat-scan.json").read_bytes())

    assert len(points) == 9
    assert points[0].bacnet_object == "analog-input,1"
    assert points[3].default is True
    assert scan.devices[0].objects[1].writable is True


@pytest.mark.parametrize(
    ("content", "message"),
    [
        (b"name,label,data_type,role,extra\nA,A,numeric,sensor,nope\n", "unsupported"),
        (b"name,label,data_type,role,default\nA,A,boolean,sensor,maybe\n", "true or false"),
        (b"name,label,data_type,role\n!!!,A,numeric,sensor\n", "usable identifier"),
    ],
)
def test_points_csv_rejects_bad_contractor_data(content: bytes, message: str) -> None:
    with pytest.raises(IntakeError, match=message):
        parse_points_csv(content)


def test_scan_json_rejects_duplicate_device_instances() -> None:
    content = (
        b'{"source":"x","devices":['
        b'{"device_instance":1,"address":"a","name":"a","objects":[]},'
        b'{"device_instance":1,"address":"b","name":"b","objects":[]}'
        b"]}"
    )
    with pytest.raises(IntakeError, match="device instances must be unique"):
        parse_bacnet_scan_json(content)


def test_contractor_xlsx_aliases_and_io_types_are_normalized() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Project point schedule"])
    sheet.append(
        [
            "Point Name",
            "Description",
            "Point Type",
            "Direction",
            "Engineering Units",
            "Device ID",
            "Object ID",
        ]
    )
    sheet.append(["ZoneTemp", "Zone temperature", "AI", "", "degF", 120012, "analog-input,1"])
    sheet.append(["FanCommand", "Supply fan", "BO", "", "", 120012, "binary-output,1"])
    content = io.BytesIO()
    workbook.save(content)

    points = parse_points_file(content.getvalue(), "contractor-points.xlsx")

    assert [point.name for point in points] == ["ZoneTemp", "FanCommand"]
    assert points[0].role == "sensor"
    assert points[1].role == "command"
    assert points[1].data_type == "boolean"
    assert points[0].bacnet_device_instance == 120012


def test_spaces_and_split_bacnet_columns_preserve_source_name() -> None:
    content = (
        b"Point Name,Description,Object Type,Object Instance,Device ID\n"
        b"AHU-1 SAT,Supply air temperature,AI,7,1001\n"
        b"AHU-1 SF-C,Supply fan command,BO,3,1001\n"
    )

    points = parse_points_csv(content)

    assert [point.name for point in points] == ["AHU_1_SAT", "AHU_1_SF_C"]
    assert points[0].source_name == "AHU-1 SAT"
    assert points[0].data_type == "numeric"
    assert points[0].role == "sensor"
    assert points[0].bacnet_object == "analog-input,7"
    assert points[1].data_type == "boolean"
    assert points[1].role == "command"
    assert points[1].bacnet_object == "binary-output,3"


def test_exact_niagara_ord_and_command_priority_are_ingested() -> None:
    content = (
        b"name,data_type,role,station_ord,write_priority\n"
        b"FanCommand,boolean,command,slot:/Config/Drivers/BacnetNetwork/Fan/Points/Cmd,8\n"
    )

    point = parse_points_csv(content)[0]

    assert point.niagara_ord == (
        "station:|slot:/Config/Drivers/BacnetNetwork/Fan/Points/Cmd"
    )
    assert point.niagara_write_priority == 8


def test_normalized_point_name_collisions_fail_closed() -> None:
    content = b"name,data_type,role\nAHU-1 SAT,numeric,sensor\nAHU 1 SAT,numeric,sensor\n"
    with pytest.raises(IntakeError, match="duplicate point identifier"):
        parse_points_csv(content)


def test_sequence_docx_is_extracted_hashed_and_classified() -> None:
    document = Document()
    document.add_heading("Variable Air Volume Terminal", level=1)
    document.add_paragraph("The VAV terminal shall modulate its damper and provide reheat.")
    content = io.BytesIO()
    document.save(content)

    parsed = parse_sequence_document(
        content.getvalue(),
        "sequence.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert "provide reheat" in parsed.text
    assert len(parsed.sha256) == 64
    assert parsed.suggested_sequence_families[0]["family"] == "G36_VAV_REHEAT"


def test_sequence_json_may_declare_an_explicit_family() -> None:
    content = json.dumps(
        {
            "family": "EXHAUST_FAN_PROOF",
            "sequence_of_operations": "Start the exhaust fan and prove status.",
        }
    ).encode()

    parsed = parse_sequence_document(content, "sequence.json", "application/json")

    assert parsed.suggested_sequence_families == [
        {
            "family": "EXHAUST_FAN_PROOF",
            "confidence": 1.0,
            "evidence": "declared in JSON",
        }
    ]


def test_fractional_bacnet_device_instance_is_rejected() -> None:
    content = b"name,data_type,role,bacnet_device_instance\nA,numeric,sensor,1.5\n"
    with pytest.raises(IntakeError, match="must be an integer"):
        parse_points_csv(content)
