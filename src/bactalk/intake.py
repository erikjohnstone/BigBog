from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import zipfile
from pathlib import Path, PurePosixPath
from typing import Any

from pydantic import BaseModel, ConfigDict, Field, ValidationError

from bactalk.domain import BacnetScan, DataType, PointRole, PointSpec

MAX_POINTS_CSV_BYTES = 2 * 1024 * 1024
MAX_POINTS_XLSX_BYTES = 10 * 1024 * 1024
MAX_SCAN_JSON_BYTES = 10 * 1024 * 1024
MAX_SEQUENCE_DOCUMENT_BYTES = 20 * 1024 * 1024
MAX_SEQUENCE_TEXT_CHARACTERS = 2_000_000
MAX_TEMPLATE_BOG_BYTES = 50 * 1024 * 1024
MAX_TEMPLATE_XML_BYTES = 100 * 1024 * 1024
MAX_POINT_ROWS = 10_000
MAX_OFFICE_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_PDF_PAGES = 500

POINT_COLUMNS = {
    "name",
    "label",
    "data_type",
    "role",
    "units",
    "default",
    "required",
    "bacnet_device_instance",
    "bacnet_object",
    "bacnet_object_type",
    "bacnet_object_instance",
    "niagara_ord",
    "niagara_write_priority",
    "brick_class",
}
POINT_COLUMN_ALIASES = {
    "point": "name",
    "point_name": "name",
    "tag": "name",
    "symbol": "name",
    "description": "label",
    "desc": "label",
    "display_name": "label",
    "type": "data_type",
    "datatype": "data_type",
    "data_type": "data_type",
    "signal_type": "data_type",
    "point_type": "data_type",
    "direction": "role",
    "function": "role",
    "io_role": "role",
    "unit": "units",
    "engineering_units": "units",
    "default_value": "default",
    "initial_value": "default",
    "device_instance": "bacnet_device_instance",
    "device_id": "bacnet_device_instance",
    "object": "bacnet_object",
    "object_id": "bacnet_object",
    "bacnet_object_id": "bacnet_object",
    "object_type": "bacnet_object_type",
    "bacnet_type": "bacnet_object_type",
    "object_instance": "bacnet_object_instance",
    "object_number": "bacnet_object_instance",
    "station_ord": "niagara_ord",
    "slot_ord": "niagara_ord",
    "niagara_path": "niagara_ord",
    "write_priority": "niagara_write_priority",
    "bacnet_write_priority": "niagara_write_priority",
    "brick": "brick_class",
}

BACNET_OBJECT_TYPES = {
    "ai": "analog-input",
    "analoginput": "analog-input",
    "ao": "analog-output",
    "analogoutput": "analog-output",
    "av": "analog-value",
    "analogvalue": "analog-value",
    "bi": "binary-input",
    "binaryinput": "binary-input",
    "bo": "binary-output",
    "binaryoutput": "binary-output",
    "bv": "binary-value",
    "binaryvalue": "binary-value",
}


class SequenceDocument(BaseModel):
    model_config = ConfigDict(extra="forbid")

    filename: str = Field(min_length=1, max_length=255)
    media_type: str = Field(min_length=1, max_length=200)
    sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    text: str = Field(min_length=1, max_length=MAX_SEQUENCE_TEXT_CHARACTERS)
    suggested_sequence_families: list[dict[str, str | float]] = Field(default_factory=list)


class IntakeError(ValueError):
    pass


def _decode_utf8(content: bytes, label: str) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise IntakeError(f"{label} must be UTF-8 encoded") from exc


def _parse_bool(value: str, *, field: str, row_number: int) -> bool:
    normalized = value.strip().lower()
    if normalized in {"true", "1", "yes", "y"}:
        return True
    if normalized in {"false", "0", "no", "n"}:
        return False
    raise IntakeError(f"points CSV row {row_number}: {field} must be true or false")


def _normal_header(value: Any) -> str:
    normalized = re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")
    return POINT_COLUMN_ALIASES.get(normalized, normalized)


def _point_symbol(value: str, *, row_number: int) -> tuple[str, str | None]:
    source = value.strip()
    symbol = re.sub(r"[^A-Za-z0-9_]+", "_", source).strip("_")
    if not symbol:
        raise IntakeError(f"points row {row_number}: point name has no usable identifier")
    if symbol[0].isdigit():
        symbol = f"P_{symbol}"
    if len(symbol) > 120:
        digest = hashlib.sha256(source.encode()).hexdigest()[:10]
        symbol = f"{symbol[:109]}_{digest}"
    return symbol, source if symbol != source else None


def _integer(value: str, *, field: str, row_number: int) -> int | None:
    if not value:
        return None
    try:
        number = float(value)
        if not number.is_integer():
            raise ValueError
        return int(number)
    except ValueError as exc:
        raise IntakeError(f"points row {row_number}: {field} must be an integer") from exc


def _bacnet_object_type(value: str, *, row_number: int) -> str:
    normalized = re.sub(r"[^a-z0-9]+", "", value.lower())
    object_type = BACNET_OBJECT_TYPES.get(normalized)
    if object_type is None:
        raise IntakeError(f"points row {row_number}: unsupported BACnet object type {value!r}")
    return object_type


def _bacnet_object(value: str, *, row_number: int) -> str:
    match = re.fullmatch(r"\s*([A-Za-z -]+)\s*[,/:]\s*([0-9]+)\s*", value)
    if match is None:
        raise IntakeError(
            f"points row {row_number}: BACnet object must look like analog-input,1 or AI:1"
        )
    object_type = _bacnet_object_type(match.group(1), row_number=row_number)
    object_instance = int(match.group(2))
    if object_instance > 4_194_302:
        raise IntakeError(f"points row {row_number}: BACnet object instance exceeds 4,194,302")
    return f"{object_type},{object_instance}"


def _data_type(value: str, *, row_number: int) -> tuple[DataType, PointRole | None]:
    normalized = re.sub(r"[^a-z0-9]+", "", value.lower())
    mapping = {
        "numeric": (DataType.NUMERIC, None),
        "number": (DataType.NUMERIC, None),
        "analog": (DataType.NUMERIC, None),
        "real": (DataType.NUMERIC, None),
        "float": (DataType.NUMERIC, None),
        "integer": (DataType.NUMERIC, None),
        "ai": (DataType.NUMERIC, PointRole.SENSOR),
        "analoginput": (DataType.NUMERIC, PointRole.SENSOR),
        "ao": (DataType.NUMERIC, PointRole.COMMAND),
        "analogoutput": (DataType.NUMERIC, PointRole.COMMAND),
        "av": (DataType.NUMERIC, PointRole.SETPOINT),
        "analogvalue": (DataType.NUMERIC, PointRole.SETPOINT),
        "boolean": (DataType.BOOLEAN, None),
        "bool": (DataType.BOOLEAN, None),
        "binary": (DataType.BOOLEAN, None),
        "digital": (DataType.BOOLEAN, None),
        "bi": (DataType.BOOLEAN, PointRole.STATUS),
        "binaryinput": (DataType.BOOLEAN, PointRole.STATUS),
        "bo": (DataType.BOOLEAN, PointRole.COMMAND),
        "binaryoutput": (DataType.BOOLEAN, PointRole.COMMAND),
        "bv": (DataType.BOOLEAN, PointRole.SETPOINT),
        "binaryvalue": (DataType.BOOLEAN, PointRole.SETPOINT),
    }
    result = mapping.get(normalized)
    if result is None:
        raise IntakeError(f"points row {row_number}: unsupported data type {value!r}")
    return result


def _point_role(value: str, *, row_number: int) -> PointRole:
    normalized = re.sub(r"[^a-z0-9]+", "", value.lower())
    mapping = {
        "sensor": PointRole.SENSOR,
        "input": PointRole.SENSOR,
        "setpoint": PointRole.SETPOINT,
        "sp": PointRole.SETPOINT,
        "command": PointRole.COMMAND,
        "cmd": PointRole.COMMAND,
        "output": PointRole.COMMAND,
        "status": PointRole.STATUS,
        "proof": PointRole.STATUS,
        "feedback": PointRole.STATUS,
        "alarm": PointRole.ALARM,
        "fault": PointRole.ALARM,
    }
    role = mapping.get(normalized)
    if role is None:
        raise IntakeError(f"points row {row_number}: unsupported role {value!r}")
    return role


def _parse_point_rows(headers: list[Any], rows: list[list[Any]]) -> list[PointSpec]:
    canonical = [_normal_header(value) for value in headers]
    duplicates = sorted({name for name in canonical if name and canonical.count(name) > 1})
    if duplicates:
        raise IntakeError("points list has duplicate normalized columns: " + ", ".join(duplicates))
    fieldnames = {name for name in canonical if name}
    unknown = sorted(fieldnames - POINT_COLUMNS)
    if unknown:
        raise IntakeError(f"points list has unsupported columns: {', '.join(unknown)}")
    if "name" not in fieldnames or not {"data_type", "bacnet_object_type"} & fieldnames:
        raise IntakeError(
            "points list requires a point name and either data type or BACnet object type column"
        )

    points: list[PointSpec] = []
    point_names: set[str] = set()
    for row_number, values in enumerate(rows, start=2):
        if len(points) >= MAX_POINT_ROWS:
            raise IntakeError(f"points list exceeds the {MAX_POINT_ROWS:,} row limit")
        row = {
            name: str(values[index]).strip()
            if index < len(values) and values[index] is not None
            else ""
            for index, name in enumerate(canonical)
            if name
        }
        if not any(row.values()):
            continue
        name, source_name = _point_symbol(row.get("name", ""), row_number=row_number)
        if name in point_names:
            raise IntakeError(
                f"points row {row_number}: duplicate point identifier after normalization: {name}"
            )
        point_names.add(name)
        type_text = row.get("data_type") or row.get("bacnet_object_type", "")
        data_type, inferred_role = _data_type(type_text, row_number=row_number)
        if row.get("bacnet_object_type"):
            object_data_type, object_role = _data_type(
                row["bacnet_object_type"], row_number=row_number
            )
            if object_data_type != data_type:
                raise IntakeError(
                    f"points row {row_number}: data type conflicts with BACnet object type"
                )
            inferred_role = inferred_role or object_role
        role = _point_role(row["role"], row_number=row_number) if row.get("role") else inferred_role
        if role is None:
            raise IntakeError(
                f"points row {row_number}: role is required when data type does not imply I/O"
            )

        default_text = row.get("default", "")
        if data_type == DataType.BOOLEAN:
            default: float | bool = (
                _parse_bool(default_text, field="default", row_number=row_number)
                if default_text
                else False
            )
        else:
            try:
                default = float(default_text) if default_text else 0.0
            except ValueError as exc:
                raise IntakeError(f"points row {row_number}: default must be numeric") from exc

        required_text = row.get("required", "")
        required = (
            _parse_bool(required_text, field="required", row_number=row_number)
            if required_text
            else True
        )
        instance = _integer(
            row.get("bacnet_device_instance", ""),
            field="bacnet_device_instance",
            row_number=row_number,
        )
        object_id = row.get("bacnet_object", "")
        object_type = row.get("bacnet_object_type", "")
        object_instance_text = row.get("bacnet_object_instance", "")
        if object_id:
            object_id = _bacnet_object(object_id, row_number=row_number)
        elif object_type or object_instance_text:
            if not object_type or not object_instance_text:
                raise IntakeError(
                    f"points row {row_number}: BACnet object type and instance must be supplied "
                    "together"
                )
            object_instance = _integer(
                object_instance_text,
                field="bacnet_object_instance",
                row_number=row_number,
            )
            if object_instance is None or not 0 <= object_instance <= 4_194_302:
                raise IntakeError(
                    f"points row {row_number}: bacnet_object_instance must be between 0 and "
                    "4,194,302"
                )
            object_id = (
                f"{_bacnet_object_type(object_type, row_number=row_number)},{object_instance}"
            )

        payload: dict[str, Any] = {
            "name": name,
            "source_name": source_name,
            "label": row.get("label") or row["name"],
            "data_type": data_type,
            "role": role,
            "units": row.get("units") or None,
            "default": default,
            "required": required,
            "bacnet_device_instance": instance,
            "bacnet_object": object_id or None,
            "niagara_ord": row.get("niagara_ord") or None,
            "niagara_write_priority": _integer(
                row.get("niagara_write_priority", ""),
                field="niagara_write_priority",
                row_number=row_number,
            ),
            "brick_class": row.get("brick_class") or None,
        }
        try:
            points.append(PointSpec.model_validate(payload))
        except ValidationError as exc:
            message = exc.errors(include_url=False)[0]["msg"]
            raise IntakeError(f"points row {row_number}: {message}") from exc
    if not points:
        raise IntakeError("points list contains no point rows")
    return points


def parse_points_csv(content: bytes) -> list[PointSpec]:
    if not content:
        raise IntakeError("points CSV is empty")
    if len(content) > MAX_POINTS_CSV_BYTES:
        raise IntakeError("points CSV exceeds the 2 MiB limit")

    reader = csv.DictReader(io.StringIO(_decode_utf8(content, "points CSV")))
    if not reader.fieldnames:
        raise IntakeError("points CSV is missing a header row")
    rows = [[raw.get(name, "") for name in reader.fieldnames] for raw in reader]
    return _parse_point_rows(list(reader.fieldnames), rows)


def _validate_office_archive(content: bytes, label: str) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > 10_000:
                raise IntakeError(f"{label} contains too many archive members")
            total = sum(item.file_size for item in infos)
            if total > MAX_OFFICE_UNCOMPRESSED_BYTES:
                raise IntakeError(f"{label} expands beyond the 100 MiB safety limit")
            for info in infos:
                path = PurePosixPath(info.filename)
                if path.is_absolute() or ".." in path.parts:
                    raise IntakeError(f"{label} contains an unsafe archive path")
                if info.compress_size and info.file_size / info.compress_size > 200:
                    raise IntakeError(f"{label} contains a suspiciously compressed member")
                if info.filename.lower().endswith(".xml"):
                    prefix = archive.read(info)[:65_536].upper()
                    if b"<!DOCTYPE" in prefix or b"<!ENTITY" in prefix:
                        raise IntakeError(f"{label} contains forbidden XML declarations")
    except zipfile.BadZipFile as exc:
        raise IntakeError(f"{label} is not a valid Office archive") from exc


def parse_points_xlsx(content: bytes) -> list[PointSpec]:
    if not content:
        raise IntakeError("points workbook is empty")
    if len(content) > MAX_POINTS_XLSX_BYTES:
        raise IntakeError("points workbook exceeds the 10 MiB limit")
    _validate_office_archive(content, "points workbook")
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as exc:
        raise IntakeError("points workbook could not be read safely") from exc
    try:
        for sheet in workbook.worksheets[:20]:
            materialized = sheet.iter_rows(values_only=True)
            leading: list[list[Any]] = []
            for _ in range(50):
                try:
                    leading.append(list(next(materialized)))
                except StopIteration:
                    break
            header_index = next(
                (
                    index
                    for index, row in enumerate(leading)
                    if "name" in {_normal_header(value) for value in row}
                    and {
                        "data_type",
                        "bacnet_object_type",
                    }
                    & {_normal_header(value) for value in row}
                ),
                None,
            )
            if header_index is None:
                continue
            rows = leading[header_index + 1 :] + [list(row) for row in materialized]
            return _parse_point_rows(leading[header_index], rows)
    finally:
        workbook.close()
    raise IntakeError("points workbook has no sheet with recognizable point headers")


def parse_points_file(content: bytes, filename: str) -> list[PointSpec]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return parse_points_csv(content)
    if suffix == ".xlsx":
        return parse_points_xlsx(content)
    raise IntakeError("points list must be a .csv or .xlsx file; macro workbooks are rejected")


def _sequence_suggestions(text: str, explicit_family: str | None) -> list[dict[str, str | float]]:
    if explicit_family:
        return [{"family": explicit_family, "confidence": 1.0, "evidence": "declared in JSON"}]
    lowered = text.lower()
    candidates = [
        (
            "G36_VAV_REHEAT",
            ("vav", "reheat"),
            "contains VAV and reheat terminology",
        ),
        (
            "AHU_DUCT_STATIC_PI",
            ("duct", "static pressure"),
            "contains duct static-pressure terminology",
        ),
        (
            "EXHAUST_FAN_PROOF",
            ("exhaust", "proof"),
            "contains exhaust fan proof terminology",
        ),
        (
            "TWO_PUMP_AVAILABILITY_SELECTOR",
            ("pump", "duty", "standby"),
            "contains pump duty/standby terminology",
        ),
        (
            "CUSTOM_AHU_SAFETY_COOLING",
            ("air handling", "supply fan", "cooling"),
            "contains AHU fan and cooling terminology",
        ),
    ]
    suggestions: list[dict[str, str | float]] = []
    for family, terms, evidence in candidates:
        hits = sum(term in lowered for term in terms)
        if hits >= 2:
            suggestions.append(
                {
                    "family": family,
                    "confidence": round(hits / len(terms), 3),
                    "evidence": evidence,
                }
            )
    return sorted(suggestions, key=lambda item: float(item["confidence"]), reverse=True)


def parse_sequence_document(
    content: bytes,
    filename: str,
    media_type: str | None = None,
) -> SequenceDocument:
    if not content:
        raise IntakeError("sequence document is empty")
    if len(content) > MAX_SEQUENCE_DOCUMENT_BYTES:
        raise IntakeError("sequence document exceeds the 20 MiB limit")
    suffix = Path(filename).suffix.lower()
    declared_family: str | None = None
    if suffix in {".txt", ".md"}:
        text = _decode_utf8(content, "sequence document")
    elif suffix == ".json":
        try:
            payload = json.loads(_decode_utf8(content, "sequence document"))
        except json.JSONDecodeError as exc:
            raise IntakeError(f"sequence JSON is invalid: {exc.msg}") from exc
        if isinstance(payload, str):
            text = payload
        elif isinstance(payload, dict):
            family = payload.get("family")
            if family is not None and not isinstance(family, str):
                raise IntakeError("sequence JSON family must be a string")
            declared_family = family
            text_value = next(
                (
                    payload[key]
                    for key in ("sequence_of_operations", "sequence", "text")
                    if isinstance(payload.get(key), str)
                ),
                None,
            )
            text = text_value or json.dumps(payload, indent=2, sort_keys=True)
        else:
            raise IntakeError("sequence JSON must contain an object or string")
    elif suffix == ".docx":
        _validate_office_archive(content, "sequence DOCX")
        from docx import Document

        try:
            document = Document(io.BytesIO(content))
        except Exception as exc:
            raise IntakeError("sequence DOCX could not be read safely") from exc
        fragments = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            fragments.extend(
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in table.rows
                if any(cell.text.strip() for cell in row.cells)
            )
        text = "\n".join(fragments)
    elif suffix == ".pdf":
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(content), strict=True)
            if reader.is_encrypted:
                raise IntakeError("encrypted sequence PDFs are not accepted")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise IntakeError(f"sequence PDF exceeds the {MAX_PDF_PAGES}-page limit")
            fragments = []
            extracted_characters = 0
            for page in reader.pages:
                fragment = page.extract_text() or ""
                fragments.append(fragment)
                extracted_characters += len(fragment)
                if extracted_characters > MAX_SEQUENCE_TEXT_CHARACTERS:
                    raise IntakeError(
                        "extracted sequence text exceeds the 2,000,000-character limit"
                    )
            text = "\n\n".join(fragments)
        except IntakeError:
            raise
        except Exception as exc:
            raise IntakeError("sequence PDF could not be parsed safely") from exc
    else:
        raise IntakeError("sequence document must be .txt, .md, .json, .docx, or .pdf")

    text = text.replace("\x00", "").strip()
    if not text:
        raise IntakeError("sequence document contains no extractable text")
    if len(text) > MAX_SEQUENCE_TEXT_CHARACTERS:
        raise IntakeError("extracted sequence text exceeds the 2,000,000-character limit")
    return SequenceDocument(
        filename=Path(filename).name,
        media_type=media_type or "application/octet-stream",
        sha256=hashlib.sha256(content).hexdigest(),
        text=text,
        suggested_sequence_families=_sequence_suggestions(text, declared_family),
    )


def parse_bacnet_scan_json(content: bytes) -> BacnetScan:
    if not content:
        raise IntakeError("BACnet scan JSON is empty")
    if len(content) > MAX_SCAN_JSON_BYTES:
        raise IntakeError("BACnet scan JSON exceeds the 10 MiB limit")
    try:
        decoded = json.loads(_decode_utf8(content, "BACnet scan JSON"))
        return BacnetScan.model_validate(decoded)
    except json.JSONDecodeError as exc:
        raise IntakeError(f"BACnet scan JSON is invalid: {exc.msg}") from exc
    except ValidationError as exc:
        message = exc.errors(include_url=False)[0]["msg"]
        raise IntakeError(f"BACnet scan JSON is invalid: {message}") from exc


def validate_template_bog(content: bytes) -> None:
    if not content:
        raise IntakeError("Niagara template is empty")
    if len(content) > MAX_TEMPLATE_BOG_BYTES:
        raise IntakeError("Niagara template exceeds the 50 MiB limit")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = archive.namelist()
            if names != ["file.xml"]:
                raise IntakeError("Niagara template must contain exactly one file.xml entry")
            info = archive.getinfo("file.xml")
            if info.file_size > MAX_TEMPLATE_XML_BYTES:
                raise IntakeError("Niagara template file.xml exceeds the 100 MiB limit")
            prefix = archive.read("file.xml", pwd=None)[:256].lstrip()
            if not prefix.startswith(b"<?xml") and not prefix.startswith(b"<"):
                raise IntakeError("Niagara template file.xml is not XML")
    except (zipfile.BadZipFile, RuntimeError) as exc:
        raise IntakeError("Niagara template is not a valid .bog ZIP archive") from exc
